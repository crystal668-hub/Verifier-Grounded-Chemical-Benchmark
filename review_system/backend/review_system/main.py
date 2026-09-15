import io, json, zipfile, hashlib, mimetypes, uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Annotated
from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi import Request, Response
from sqlalchemy import delete, func, or_, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session as DBSession
from .config import ADMIN_PASSWORD, ADMIN_USER, ALLOWED_ORIGINS, AUTO_CREATE_DB, COOKIE_SECURE, LOGIN_FAILURE_LIMIT, LOGIN_WINDOW_MINUTES, MAX_ATTACHMENT_BYTES, DATA_DIR, validate_production_config
from .db import SessionLocal, init_db
from .models import Attachment, Comment, Draft, DraftRevision, LoginAttempt, PasswordChangeEvent, ReviewEvent, Session as UserSession, SnapshotTask, SnapshotTrack, SourceSnapshot, User, SharedFile
from .schemas import CommentIn, DecisionIn, DraftIn, LoginIn, PasswordChangeIn, PasswordResetIn, RegisterIn, UserIn, UserUpdateIn
from .security import client_ip, create_session, current_user, hash_password, require_developer, revoke_user_sessions, session_token_hash, verify_password
from .source import load_catalog, schema_view, scoring_view, split_attachments
from verifier_grounded_benchmark import load_track
import markdown, bleach
from openpyxl import load_workbook

app=FastAPI(title="VGB Task Review API", version="0.1.0")
app.add_middleware(CORSMiddleware, allow_origins=ALLOWED_ORIGINS, allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

BEIJING_TZ = timezone(timedelta(hours=8))

def beijing_isoformat(value: datetime | None) -> str | None:
    if value is None:
        return None
    utc_value = value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)
    return utc_value.astimezone(BEIJING_TZ).isoformat()

@app.middleware("http")
async def security_headers(request: Request, call_next):
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "same-origin"
    if COOKIE_SECURE:
        response.headers["Strict-Transport-Security"] = "max-age=31536000"
    return response

def db():
    session=SessionLocal()
    try: yield session
    finally: session.close()

def auth(request: Request, database: DBSession = Depends(db)) -> User: return current_user(request, database)
def developer(user: User = Depends(auth)) -> User: return require_developer(user)

@app.on_event("startup")
def startup():
    if AUTO_CREATE_DB:
        init_db()
    with SessionLocal() as database:
        database_is_empty = database.scalar(select(User.id).limit(1)) is None
        validate_production_config(database_is_empty=database_is_empty)
        if database_is_empty:
            database.add(User(username=ADMIN_USER,password_hash=hash_password(ADMIN_PASSWORD or ""),role="developer")); database.commit()
        if database.scalar(select(SourceSnapshot).where(SourceSnapshot.status=="active")) is None:
            sync_catalog(database)

def sync_catalog(database: DBSession) -> SourceSnapshot:
    catalog=load_catalog(); existing=database.scalar(select(SourceSnapshot).where(SourceSnapshot.fingerprint==catalog["fingerprint"]))
    if existing: return existing
    snapshot=SourceSnapshot(commit=catalog["commit"], fingerprint=catalog["fingerprint"], status="active")
    database.add(snapshot); database.flush()
    for track in catalog["tracks"]:
        st=SnapshotTrack(snapshot_id=snapshot.id,name=track["name"],version=track["version"],display_name=track["display_name"]); database.add(st); database.flush()
        for task in track["tasks"]:
            database.add(SnapshotTask(track_id=st.id,task_id=task["task_id"],version=task["version"],fingerprint=task["fingerprint"],data_json=json.dumps(task["data"],ensure_ascii=False),view_json=json.dumps(task["view"],ensure_ascii=False),scoring_json=json.dumps(task["scoring"],ensure_ascii=False),schema_json=json.dumps(task["schema"],ensure_ascii=False),attachments_json=json.dumps(task["attachments"],ensure_ascii=False)))
    database.commit(); return snapshot

def active_snapshot(database: DBSession) -> SourceSnapshot:
    snapshot=database.scalar(select(SourceSnapshot).where(SourceSnapshot.status=="active").order_by(SourceSnapshot.id.desc()))
    if not snapshot: raise HTTPException(503,"题库尚未同步")
    return snapshot

@app.get("/healthz")
def healthz(database: DBSession = Depends(db)):
    database.execute(text("SELECT 1"))
    if database.scalar(select(SourceSnapshot.id).where(SourceSnapshot.status == "active").limit(1)) is None:
        raise HTTPException(503, "not ready")
    return {"status": "ok"}

def login_rate_limited(database: DBSession, username: str, source_ip: str) -> bool:
    cutoff = datetime.now(timezone.utc) - timedelta(minutes=LOGIN_WINDOW_MINUTES)
    count = database.scalar(select(func.count(LoginAttempt.id)).where(LoginAttempt.successful.is_(False), LoginAttempt.created_at >= cutoff, or_(LoginAttempt.username == username, LoginAttempt.source_ip == source_ip))) or 0
    return count >= LOGIN_FAILURE_LIMIT

@app.post("/api/v1/auth/login")
def login(payload: LoginIn, request: Request, response: Response, database: DBSession = Depends(db)):
    source_ip = client_ip(request)
    if login_rate_limited(database, payload.username, source_ip):
        raise HTTPException(429, "登录尝试过多，请稍后重试")
    user=database.scalar(select(User).where(User.username==payload.username))
    successful = bool(user and user.active and verify_password(payload.password,user.password_hash))
    database.add(LoginAttempt(username=payload.username, source_ip=source_ip, successful=successful))
    database.commit()
    if not successful: raise HTTPException(401,"用户名或密码错误")
    token,csrf=create_session(database,user); response.set_cookie("review_session",token,httponly=True,samesite="lax",secure=COOKIE_SECURE,max_age=604800,path="/"); return {"user":{"id":user.id,"username":user.username,"role":user.role},"csrf_token":csrf}

@app.post("/api/v1/auth/register", status_code=status.HTTP_201_CREATED)
def register(payload: RegisterIn, request: Request, response: Response, database: DBSession = Depends(db)):
    username = payload.username.strip()
    if not username:
        raise HTTPException(422, "用户名不能为空")
    user = User(username=username, password_hash=hash_password(payload.password), role="collaborator", active=True)
    database.add(user)
    try:
        database.commit()
    except IntegrityError:
        database.rollback()
        raise HTTPException(409, "用户名已存在") from None
    database.refresh(user)
    token, csrf = create_session(database, user)
    response.set_cookie("review_session", token, httponly=True, samesite="lax", secure=COOKIE_SECURE, max_age=604800, path="/")
    return {"user": {"id": user.id, "username": user.username, "role": user.role}, "csrf_token": csrf}

@app.post("/api/v1/auth/logout")
def logout(response: Response, request: Request, database: DBSession = Depends(db), user: User = Depends(auth)):
    token = request.cookies.get("review_session")
    if token:
        session = database.get(UserSession, session_token_hash(token))
        if session:
            database.delete(session); database.commit()
    response.delete_cookie("review_session", path="/", secure=COOKIE_SECURE, httponly=True, samesite="lax"); return {"ok":True}

@app.get("/api/v1/auth/me")
def me(user: User = Depends(auth)): return {"id":user.id,"username":user.username,"role":user.role}

@app.post("/api/v1/auth/password")
def change_password(payload: PasswordChangeIn, database: DBSession=Depends(db), user: User=Depends(auth), idempotency_key: Annotated[str | None, Header(alias="Idempotency-Key")]=None):
    request_id=(idempotency_key or "").strip()
    if len(request_id)>80: raise HTTPException(422,"幂等请求标识过长")
    if request_id and database.scalar(select(PasswordChangeEvent).where(PasswordChangeEvent.user_id==user.id,PasswordChangeEvent.request_id==request_id)):
        return {"ok":True,"changed":False,"idempotent_replay":True}
    if not verify_password(payload.current_password,user.password_hash):
        if verify_password(payload.new_password,user.password_hash):
            return {"ok":True,"changed":False,"idempotent_replay":True}
        raise HTTPException(400,"当前密码错误")
    if verify_password(payload.new_password,user.password_hash):
        return {"ok":True,"changed":False,"idempotent_replay":False}
    user.password_hash=hash_password(payload.new_password)
    database.add(PasswordChangeEvent(user_id=user.id,request_id=request_id or uuid.uuid4().hex))
    try:
        database.commit()
    except IntegrityError:
        database.rollback()
        if request_id and database.scalar(select(PasswordChangeEvent).where(PasswordChangeEvent.user_id==user.id,PasswordChangeEvent.request_id==request_id)):
            return {"ok":True,"changed":False,"idempotent_replay":True}
        raise
    return {"ok":True,"changed":True,"idempotent_replay":False}

@app.get("/api/v1/users")
def users(database: DBSession=Depends(db), user: User=Depends(developer)):
    last_change=select(func.max(PasswordChangeEvent.created_at)).where(PasswordChangeEvent.user_id==User.id).correlate(User).scalar_subquery()
    rows=database.execute(select(User,last_change).order_by(User.username)).all()
    return [{"id":item.id,"username":item.username,"role":item.role,"active":item.active,"created_at":beijing_isoformat(item.created_at),"password_changed_at":beijing_isoformat(changed_at)} for item,changed_at in rows]

@app.post("/api/v1/users")
def create_user(payload: UserIn, database: DBSession=Depends(db), user: User=Depends(developer)):
    if database.scalar(select(User).where(User.username==payload.username)):
        raise HTTPException(409,"用户名已存在")
    row=User(username=payload.username,password_hash=hash_password(payload.password),role=payload.role); database.add(row); database.commit(); database.refresh(row); return {"id":row.id,"username":row.username,"role":row.role}

def ensure_another_developer(database: DBSession, target: User, *, active: bool | None = None, role: str | None = None) -> None:
    removes_developer = target.role == "developer" and (active is False or role == "collaborator")
    if removes_developer:
        count = database.scalar(select(func.count(User.id)).where(User.role == "developer", User.active.is_(True))) or 0
        if count <= 1:
            raise HTTPException(409, "不能停用或降级最后一个开发者")

@app.patch("/api/v1/users/{user_id}")
def update_user(user_id: int, payload: UserUpdateIn, database: DBSession=Depends(db), actor: User=Depends(developer)):
    target = database.get(User, user_id)
    if not target: raise HTTPException(404, "用户不存在")
    ensure_another_developer(database, target, active=payload.active, role=payload.role)
    if payload.active is not None: target.active = payload.active
    if payload.role is not None: target.role = payload.role
    if not target.active: revoke_user_sessions(database, target.id)
    database.commit(); return {"id":target.id,"username":target.username,"role":target.role,"active":target.active}

@app.post("/api/v1/users/{user_id}/password")
def reset_password(user_id: int, payload: PasswordResetIn, database: DBSession=Depends(db), actor: User=Depends(developer)):
    target = database.get(User, user_id)
    if not target: raise HTTPException(404, "用户不存在")
    target.password_hash = hash_password(payload.new_password); revoke_user_sessions(database, target.id); database.commit()
    return {"ok": True, "sessions_revoked": True}

@app.delete("/api/v1/users/{user_id}/sessions")
def revoke_sessions(user_id: int, database: DBSession=Depends(db), actor: User=Depends(developer)):
    if not database.get(User, user_id): raise HTTPException(404, "用户不存在")
    revoke_user_sessions(database, user_id); database.commit(); return {"ok": True}

@app.get("/api/v1/tracks")
def tracks(database: DBSession=Depends(db), user: User=Depends(auth)):
    snapshot=active_snapshot(database); return {"snapshot":{"id":snapshot.id,"commit":snapshot.commit,"fingerprint":snapshot.fingerprint},"tracks":[{"name":x.name,"version":x.version,"display_name":x.display_name,"task_count":len(x.tasks)} for x in snapshot.tracks]}

@app.get("/api/v1/tracks/{track}/tasks")
def task_list(track: str, q: str|None=None, status: str|None=None, database: DBSession=Depends(db), user: User=Depends(auth)):
    snapshot=active_snapshot(database); item=next((x for x in snapshot.tracks if x.name==track),None)
    if not item: raise HTTPException(404,"track 不存在")
    tasks=[x for x in item.tasks if (not q or q.lower() in x.task_id.lower() or q.lower() in json.loads(x.view_json).get("prompt","").lower()) and (not status or x.review_status==status)]
    return {"track":track,"tasks":[{"task_id":x.task_id,"version":x.version,"status":x.review_status,"summary":json.loads(x.view_json).get("object_type")} for x in tasks]}

def find_task(track: str, task_id: str, database: DBSession) -> SnapshotTask:
    snapshot=active_snapshot(database); item=next((x for x in snapshot.tracks if x.name==track),None)
    task=next((x for x in item.tasks if x.task_id==task_id),None) if item else None
    if not task: raise HTTPException(404,"题目不存在")
    return task

@app.get("/api/v1/tracks/{track}/tasks/{task_id}")
def task_detail(track: str, task_id: str, database: DBSession=Depends(db), user: User=Depends(auth)):
    task=find_task(track,task_id,database); view=json.loads(task.view_json); view["task_id"]=task.task_id; view["review_status"]=task.review_status; view["source_fingerprint"]=task.fingerprint; view["attachments"]=[{k:v for k,v in a.items() if k not in {"content"}} for a in json.loads(task.attachments_json)]; return view

@app.get("/api/v1/tracks/{track}/tasks/{task_id}/scoring")
def task_scoring(track: str, task_id: str, database: DBSession=Depends(db), user: User=Depends(auth)):
    task = find_task(track, task_id, database)
    stored = json.loads(task.scoring_json)
    # Older snapshots predate profile-aware scoring displays. Rebuild this read-only
    # view from the immutable task payload so existing databases gain the richer UI
    # without mutating or replacing their audit snapshot.
    if not stored.get("rules") or any(not rule.get("profile") or not rule.get("full_score_region") or rule.get("score_range", {}).get("kind") == "得分区间" for rule in stored.get("rules", [])):
        raw = json.loads(task.data_json)
        stored = scoring_view(raw, load_track(track)._task_pack.scoring_profiles)
    return stored
@app.get("/api/v1/tracks/{track}/tasks/{task_id}/schema")
def task_schema(track: str, task_id: str, database: DBSession=Depends(db), user: User=Depends(auth)):
    task = find_task(track, task_id, database)
    stored = json.loads(task.schema_json)
    if any(isinstance(item.get("value"), str) and len(item["value"].strip()) >= 80 for item in stored.get("input_objects", [])):
        view, _attachments = split_attachments(json.loads(task.data_json))
        stored = schema_view(view)
    return stored
@app.get("/api/v1/tracks/{track}/tasks/{task_id}/attachments")
def task_attachments(track: str, task_id: str, database: DBSession=Depends(db), user: User=Depends(auth)): return {"attachments":[{k:v for k,v in a.items() if k not in {"content"}} for a in json.loads(find_task(track,task_id,database).attachments_json)]}

@app.get("/api/v1/tracks/{track}/tasks/{task_id}/attachments/{name}")
def task_attachment(track: str, task_id: str, name: str, database: DBSession=Depends(db), user: User=Depends(auth)):
    item=next((x for x in json.loads(find_task(track,task_id,database).attachments_json) if x.get("name")==name),None)
    if not item: raise HTTPException(404,"附件不存在")
    return Response(item.get("content","") if isinstance(item.get("content"),str) else "", media_type=item.get("media_type","text/plain"), headers={"Content-Disposition": f'inline; filename="{name}"'})

ALLOWED_SHARED = {".md": "text/markdown", ".pdf": "application/pdf", ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}

def shared_preview(data: bytes, suffix: str) -> str:
    if suffix == ".md":
        html = markdown.markdown(data.decode("utf-8", errors="replace"), extensions=["tables", "fenced_code"])
        return bleach.clean(html, tags=["p","br","h1","h2","h3","h4","ul","ol","li","strong","em","code","pre","table","thead","tbody","tr","th","td","blockquote"], attributes={}, strip=True)
    if suffix == ".xlsx":
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = ["<div class='xlsx-preview'>"]
        for ws in wb.worksheets:
            out.append(f"<h3>{bleach.clean(ws.title)}</h3><table><tbody>")
            for idx, row in enumerate(ws.iter_rows(values_only=True)):
                if idx >= 1000: break
                out.append("<tr>" + "".join(f"<td>{bleach.clean(str(v)) if v is not None else ''}</td>" for v in row[:50]) + "</tr>")
            out.append("</tbody></table>")
        out.append("</div>"); return "".join(out)
    return ""

@app.get("/api/v1/shared-files")
def shared_files(database: DBSession=Depends(db), user: User=Depends(auth), q: str|None=None, media_type: str|None=None):
    rows = database.scalars(select(SharedFile).order_by(SharedFile.created_at.desc())).all()
    if q: rows = [r for r in rows if q.lower() in r.name.lower()]
    if media_type: rows = [r for r in rows if r.media_type == media_type]
    return [{"id":r.id,"name":r.name,"media_type":r.media_type,"size":r.size,"owner":r.owner.username,"track":r.track,"task_id":r.task_id,"created_at":(r.created_at if r.created_at.tzinfo else r.created_at.replace(tzinfo=timezone.utc)).isoformat(),"preview":bool(r.preview_html),"preview_error":r.preview_error,"can_delete":r.owner_id==user.id or user.role=="developer"} for r in rows]

@app.post("/api/v1/shared-files", status_code=201)
async def upload_shared_file(file: UploadFile=File(...), track: str|None=Form(None), task_id: str|None=Form(None), database: DBSession=Depends(db), user: User=Depends(auth)):
    name = Path(file.filename or "").name; suffix = Path(name).suffix.lower()
    if not name or suffix not in ALLOWED_SHARED: raise HTTPException(415, "仅支持 .md、.pdf、.xlsx 文件")
    data = await file.read(MAX_ATTACHMENT_BYTES + 1)
    if not data or len(data) > MAX_ATTACHMENT_BYTES: raise HTTPException(413, "文件为空或超过 20 MB 限制")
    if suffix == ".pdf" and not data.startswith(b"%PDF"): raise HTTPException(415, "PDF 文件内容无效")
    digest = hashlib.sha256(data).hexdigest(); existing = database.scalar(select(SharedFile).where(SharedFile.sha256==digest))
    if existing: return {"id": existing.id, "duplicate": True}
    target = DATA_DIR / "shared-files"; target.mkdir(parents=True, exist_ok=True); path = target / f"{uuid.uuid4().hex}{suffix}"; path.write_bytes(data)
    row = SharedFile(name=name, media_type=ALLOWED_SHARED[suffix], size=len(data), sha256=digest, storage_path=str(path), owner_id=user.id, track=track, task_id=task_id)
    try: row.preview_html = shared_preview(data, suffix)
    except Exception as exc: row.preview_error = str(exc)
    database.add(row); database.commit(); database.refresh(row)
    return {"id": row.id, "name": row.name}

@app.get("/api/v1/shared-files/{file_id}")
def shared_file_detail(file_id: int, database: DBSession=Depends(db), user: User=Depends(auth)):
    row = database.get(SharedFile, file_id)
    if not row: raise HTTPException(404, "资料不存在")
    return {"id":row.id,"name":row.name,"media_type":row.media_type,"size":row.size,"owner":row.owner.username,"track":row.track,"task_id":row.task_id,"preview_html":row.preview_html,"preview_error":row.preview_error,"can_delete":row.owner_id==user.id or user.role=="developer"}

@app.get("/api/v1/shared-files/{file_id}/content")
def shared_file_content(file_id: int, database: DBSession=Depends(db), user: User=Depends(auth)):
    row = database.get(SharedFile, file_id)
    if not row: raise HTTPException(404, "资料不存在")
    if row.media_type == "application/pdf": return Response(content=Path(row.storage_path).read_bytes(), media_type=row.media_type, headers={"Content-Disposition": f'inline; filename="{row.name}"'})
    return Response(row.preview_html or "", media_type="text/html")

@app.get("/api/v1/shared-files/{file_id}/download")
def shared_file_download(file_id: int, database: DBSession=Depends(db), user: User=Depends(auth)):
    row = database.get(SharedFile, file_id)
    if not row: raise HTTPException(404, "资料不存在")
    return StreamingResponse(open(row.storage_path, "rb"), media_type=row.media_type, headers={"Content-Disposition": f'attachment; filename="{row.name}"'})

@app.delete("/api/v1/shared-files/{file_id}")
def delete_shared_file(file_id: int, database: DBSession=Depends(db), user: User=Depends(auth)):
    row = database.get(SharedFile, file_id)
    if not row: raise HTTPException(404, "资料不存在")
    if row.owner_id != user.id and user.role != "developer": raise HTTPException(403, "只能删除自己上传的资料")
    Path(row.storage_path).unlink(missing_ok=True); database.delete(row); database.commit(); return {"ok": True}

@app.get("/api/v1/tasks/{task_id}/comments")
def comments(task_id: str, database: DBSession=Depends(db), user: User=Depends(auth)):
    rows=database.scalars(select(Comment).where(Comment.target_type=="task",Comment.target_id==task_id).order_by(Comment.created_at)).all(); return [{"id":x.id,"body":x.body,"author":x.author.username,"resolved":x.resolved,"parent_id":x.parent_id,"created_at":x.created_at,"can_delete":x.author_id==user.id or user.role=="developer"} for x in rows]

@app.post("/api/v1/tasks/{task_id}/comments")
def add_comment(task_id: str, payload: CommentIn, database: DBSession=Depends(db), user: User=Depends(auth)):
    row=Comment(target_type="task",target_id=task_id,author_id=user.id,body=payload.body,parent_id=payload.parent_id); database.add(row); database.commit(); database.refresh(row); return {"id":row.id,"body":row.body}

@app.delete("/api/v1/comments/{comment_id}")
def delete_comment(comment_id: int, database: DBSession=Depends(db), user: User=Depends(auth)):
    row=database.get(Comment,comment_id)
    if not row: raise HTTPException(404,"评论不存在")
    if row.author_id!=user.id and user.role!="developer": raise HTTPException(403,"只能删除自己的评论")
    for child in database.scalars(select(Comment).where(Comment.parent_id==row.id)).all():
        child.parent_id=None
    database.delete(row); database.commit(); return {"ok":True,"id":comment_id}

@app.post("/api/v1/comments/{comment_id}/resolve")
def resolve_comment(comment_id: int, database: DBSession=Depends(db), user: User=Depends(auth)):
    row=database.get(Comment,comment_id)
    if not row: raise HTTPException(404,"批注不存在")
    row.resolved=True; database.commit(); return {"ok":True}

@app.post("/api/v1/tasks/{task_id}/review")
def review_task(task_id: str, payload: DecisionIn, database: DBSession=Depends(db), user: User=Depends(developer)):
    task = database.scalar(select(SnapshotTask).where(SnapshotTask.task_id == task_id))
    if not task: raise HTTPException(404, "题目不存在")
    previous = task.review_status; task.review_status = payload.status
    database.add(ReviewEvent(target_type="task", target_id=task_id, from_status=previous, to_status=payload.status, note=payload.note, actor_id=user.id)); database.commit()
    return {"task_id": task_id, "status": task.review_status}

@app.post("/api/v1/source-sync")
def source_sync(database: DBSession=Depends(db), user: User=Depends(developer)):
    try: snapshot=sync_catalog(database); return {"ok":True,"snapshot_id":snapshot.id,"fingerprint":snapshot.fingerprint}
    except Exception as exc: raise HTTPException(422,f"同步失败：{exc}") from exc

@app.post("/api/v1/drafts")
def create_draft(payload: DraftIn, database: DBSession=Depends(db), user: User=Depends(auth)):
    row=Draft(owner_id=user.id,track=payload.track,title=payload.title,current_revision=1); database.add(row); database.flush(); database.add(DraftRevision(draft_id=row.id,revision=1,payload_json=json.dumps(payload.payload,ensure_ascii=False),created_by=user.id)); database.commit(); return {"id":row.id,"status":row.status,"revision":1}

@app.get("/api/v1/drafts")
def drafts(database: DBSession=Depends(db), user: User=Depends(auth)):
    rows=database.scalars(select(Draft).order_by(Draft.updated_at.desc())).all(); rows=[x for x in rows if x.owner_id==user.id or user.role=="developer" or x.status not in {"draft"}]; return [{"id":x.id,"track":x.track,"title":x.title,"status":x.status,"revision":x.current_revision} for x in rows]

@app.post("/api/v1/drafts/{draft_id}/submit")
def submit_draft(draft_id:int,database:DBSession=Depends(db),user:User=Depends(auth)):
    row=database.get(Draft,draft_id)
    if not row or (row.owner_id!=user.id and user.role!="developer"): raise HTTPException(404,"草稿不存在")
    row.status="pending"; database.commit(); return {"id":row.id,"status":row.status}

@app.post("/api/v1/drafts/{draft_id}/validate")
def validate_draft(draft_id:int,database:DBSession=Depends(db),user:User=Depends(auth)):
    row=database.get(Draft,draft_id)
    if not row or (row.owner_id!=user.id and user.role!="developer"): raise HTTPException(404,"草稿不存在")
    revision=database.scalar(select(DraftRevision).where(DraftRevision.draft_id==row.id,DraftRevision.revision==row.current_revision))
    try:
        payload=json.loads(revision.payload_json) if revision else {}
        errors=[]
        for key in ("task_id","prompt","answer_schema"):
            if not payload.get(key): errors.append(f"缺少 {key}")
        return {"valid":not errors,"errors":errors}
    except Exception as exc:
        return {"valid":False,"errors":[str(exc)]}

@app.post("/api/v1/drafts/{draft_id}/decision")
def draft_decision(draft_id:int,payload:DecisionIn,database:DBSession=Depends(db),user:User=Depends(developer)):
    row=database.get(Draft,draft_id)
    if not row: raise HTTPException(404,"草稿不存在")
    previous=row.status; row.status=payload.status; row.approved_revision=row.current_revision if payload.status=="approved" else None; database.add(ReviewEvent(target_type="draft",target_id=str(draft_id),from_status=previous,to_status=payload.status,note=payload.note,actor_id=user.id)); database.commit(); return {"id":row.id,"status":row.status}

@app.get("/api/v1/drafts/{draft_id}/export")
def export_draft(draft_id:int,database:DBSession=Depends(db),user:User=Depends(developer)):
    row=database.get(Draft,draft_id)
    if not row or row.status!="approved": raise HTTPException(409,"仅已批准草稿可导出")
    revision=database.scalar(select(DraftRevision).where(DraftRevision.draft_id==row.id,DraftRevision.revision==row.approved_revision)); payload=revision.payload_json if revision else "{}"
    content=io.BytesIO();
    with zipfile.ZipFile(content,"w",zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("draft.json",payload); archive.writestr("README.txt","由 VGB 题目审核系统导出，请开发者补全并合入题库。\n")
    content.seek(0); return StreamingResponse(content,media_type="application/zip",headers={"Content-Disposition":f'attachment; filename="draft-{draft_id}.zip"'})

_frontend_dist = Path(__file__).resolve().parents[2] / "frontend" / "dist"
if _frontend_dist.is_dir():
    app.mount("/", StaticFiles(directory=_frontend_dist, html=True), name="frontend")
